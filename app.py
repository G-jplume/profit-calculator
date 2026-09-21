"""
跨境电商新品利润核算工具 - Streamlit Application
Based on the original Excel profit calculation workbook.
Supports AMZ FBM, AMZ FBA, and Walmart WF business lines.
"""
import io
import plotly.graph_objects as go
import plotly.express as px
import pandas as pd
import streamlit as st
from openpyxl import load_workbook

from calculator import (
    ProductInput,
    WAREHOUSES,
    calc_amz_fbm_profit,
    calc_amz_fba_profit,
    calc_wf_profit,
    calc_weight_lbs,
    cm_to_inch,
)


def parse_excel_template(uploaded_file) -> dict:
    """从上传的 Excel 利润核算表中提取产品信息。
    自动扫描各计算 Sheet，查找品名、产地、售价、采购成本、包装尺寸等关键字段。
    """
    wb = load_workbook(uploaded_file, data_only=True)
    result = {}

    # 按优先级扫描计算表
    target_sheets = ["利润核算AMZ+WM", "利润测算FBA-中国工厂", "利润核算WF"]
    for sname in target_sheets:
        if sname not in wb.sheetnames:
            continue
        ws = wb[sname]

        # 扫描前30行，提取关键字段
        for row in ws.iter_rows(min_row=1, max_row=30, values_only=False):
            for cell in row:
                val = str(cell.value).strip() if cell.value else ""
                if not val:
                    continue
                val_lower = val.lower()

                # 品名
                if val_lower in ("品名", "产品名称", "product name") and not result.get("name"):
                    next_cell = ws.cell(row=cell.row, column=cell.column + 1)
                    if next_cell.value:
                        result["name"] = str(next_cell.value).strip()

                # 工厂产地
                if val_lower in ("产地", "工厂所在地", "工厂", "origin") and not result.get("origin"):
                    for offset in range(1, 4):
                        next_cell = ws.cell(row=cell.row, column=cell.column + offset)
                        if next_cell.value and str(next_cell.value).strip() in ("中国", "越南/泰国", "越南", "泰国", "马来西亚"):
                            origin_val = str(next_cell.value).strip()
                            if origin_val in ("越南", "泰国"):
                                origin_val = "越南/泰国"
                            result["origin"] = origin_val
                            break

                # 售价
                if val_lower in ("售价", "售价($)", "price", "售价($)") and not result.get("price"):
                    for offset in range(1, 4):
                        next_cell = ws.cell(row=cell.row, column=cell.column + offset)
                        if isinstance(next_cell.value, (int, float)) and next_cell.value > 0:
                            result["price"] = float(next_cell.value)
                            break

                # 采购成本
                if val_lower in ("采购成本", "采购价", "商品成本", "purchase") and not result.get("purchase"):
                    for offset in range(1, 4):
                        next_cell = ws.cell(row=cell.row, column=cell.column + offset)
                        if isinstance(next_cell.value, (int, float)) and next_cell.value > 0:
                            result["purchase"] = float(next_cell.value)
                            break

                # 汇率
                if val_lower in ("汇率", "exchange rate") and not result.get("rate"):
                    for offset in range(1, 4):
                        next_cell = ws.cell(row=cell.row, column=cell.column + offset)
                        if isinstance(next_cell.value, (int, float)) and next_cell.value > 0:
                            result["rate"] = float(next_cell.value)
                            break

                # 箱数
                if val_lower in ("箱数", "box count") and not result.get("box_count"):
                    for offset in range(1, 4):
                        next_cell = ws.cell(row=cell.row, column=cell.column + offset)
                        if isinstance(next_cell.value, (int, float)) and next_cell.value > 0:
                            result["box_count"] = int(next_cell.value)
                            break

                # 配送区间
                if val_lower in ("zone", "配送区间", "配送区域") and not result.get("zone"):
                    for offset in range(1, 4):
                        next_cell = ws.cell(row=cell.row, column=cell.column + offset)
                        if isinstance(next_cell.value, (int, float)) and next_cell.value in (3, 4, 5):
                            result["zone"] = int(next_cell.value)
                            break

                # 包装尺寸 - 长
                if val_lower in ("长(cm)", "长", "length", "长度") and not result.get("length"):
                    for offset in range(1, 4):
                        next_cell = ws.cell(row=cell.row, column=cell.column + offset)
                        if isinstance(next_cell.value, (int, float)) and next_cell.value > 0:
                            result["length"] = float(next_cell.value)
                            break

                # 包装尺寸 - 宽
                if val_lower in ("宽(cm)", "宽", "width") and not result.get("width"):
                    for offset in range(1, 4):
                        next_cell = ws.cell(row=cell.row, column=cell.column + offset)
                        if isinstance(next_cell.value, (int, float)) and next_cell.value > 0:
                            result["width"] = float(next_cell.value)
                            break

                # 包装尺寸 - 高
                if val_lower in ("高(cm)", "高", "height") and not result.get("height"):
                    for offset in range(1, 4):
                        next_cell = ws.cell(row=cell.row, column=cell.column + offset)
                        if isinstance(next_cell.value, (int, float)) and next_cell.value > 0:
                            result["height"] = float(next_cell.value)
                            break

                # 重量
                if val_lower in ("重量(kg)", "重量", "weight", "毛重(kg)") and not result.get("weight"):
                    for offset in range(1, 4):
                        next_cell = ws.cell(row=cell.row, column=cell.column + offset)
                        if isinstance(next_cell.value, (int, float)) and next_cell.value > 0:
                            result["weight"] = float(next_cell.value)
                            break

        if result:
            break

    # 设置默认值
    defaults = {
        "name": "导入产品",
        "origin": "中国",
        "price": 159.99,
        "purchase": 450.0,
        "rate": 7.0,
        "box_count": 1,
        "zone": 5,
        "length": 170.0,
        "width": 50.0,
        "height": 12.0,
        "weight": 36.0,
    }
    for k, v in defaults.items():
        if k not in result:
            result[k] = v

    return result

st.set_page_config(
    page_title="新品利润核算工具",
    page_icon="📊",
    layout="wide",
    initial_sidebar_state="expanded",
)

st.markdown(
    """
    <style>
    .main-header {
        background: linear-gradient(90deg, #6c4ee0 0%, #9d80f0 100%);
        padding: 20px 24px;
        border-radius: 12px;
        margin-bottom: 24px;
    }
    .main-header h1 { color: white; margin: 0; font-size: 24px; }
    .main-header p { color: rgba(255,255,255,0.85); margin: 4px 0 0; font-size: 14px; }
    .metric-card {
        background: #f7f8fa;
        border-radius: 8px;
        padding: 16px;
        border: 1px solid #e2e5ea;
    }
    .metric-card h3 { margin: 0 0 4px; font-size: 13px; color: #6b7384; text-transform: uppercase; letter-spacing: 0.05em; }
    .metric-card .value { font-size: 24px; font-weight: 700; }
    .metric-card .value.positive { color: #2ea66c; }
    .metric-card .value.negative { color: #e74c3c; }
    .metric-card .value.neutral { color: #6c4ee0; }
    </style>
    """,
    unsafe_allow_html=True,
)

st.markdown(
    """
    <div class="main-header">
        <h1>📊 跨境电商新品利润核算工具</h1>
        <p>支持 AMZ FBM / AMZ FBA / Walmart WF 三大业务线 · 6大海外仓自动比价 · 淡旺季双期对比</p>
    </div>
    """,
    unsafe_allow_html=True,
)


# ============================================================
# 默认值
# ============================================================
defaults = {
    "name": "示例产品",
    "origin": "中国",
    "zone": 5,
    "box_count": 1,
    "price": 159.99,
    "purchase": 450.0,
    "rate": 7.0,
    "length": 170.0,
    "width": 50.0,
    "height": 12.0,
    "weight": 36.0,
}

# ============================================================
# Sidebar: Excel 模板导入
# ============================================================
st.sidebar.markdown("### 导入 Excel 模板")
st.sidebar.caption("上传利润核算表 .xlsx，自动提取产品信息")

uploaded_file = st.sidebar.file_uploader(
    "选择 Excel 文件",
    type=["xlsx", "xls"],
    key="excel_upload",
    help="支持原始利润核算表格式，自动提取品名/产地/售价/采购成本/包装尺寸等字段",
)

if uploaded_file is not None:
    try:
        parsed = parse_excel_template(uploaded_file)
        defaults.update(parsed)
        st.sidebar.success(f"已导入: {defaults.get('name', '未知产品')}")
        with st.sidebar.expander("查看提取的字段"):
            display_fields = {
                "品名": defaults.get("name"),
                "产地": defaults.get("origin"),
                "售价": defaults.get("price"),
                "采购成本": defaults.get("purchase"),
                "汇率": defaults.get("rate"),
                "箱数": defaults.get("box_count"),
                "配送区间": defaults.get("zone"),
                "长(cm)": defaults.get("length"),
                "宽(cm)": defaults.get("width"),
                "高(cm)": defaults.get("height"),
                "重量(kg)": defaults.get("weight"),
            }
            for label, val in display_fields.items():
                st.sidebar.text(f"{label}: {val}")
    except Exception as e:
        st.sidebar.error(f"导入失败: {e}")

st.sidebar.markdown("---")

# ============================================================
# Sidebar: 手动填写产品信息
# ============================================================
st.sidebar.markdown("### 产品信息录入")
st.sidebar.caption("导入模板后自动填充，也可手动修改")

product_name = st.sidebar.text_input("品名", value=defaults["name"], key="name")

origin = st.sidebar.selectbox(
    "工厂所在地",
    ["中国", "越南/泰国", "马来西亚"],
    index=["中国", "越南/泰国", "马来西亚"].index(defaults["origin"]) if defaults["origin"] in ["中国", "越南/泰国", "马来西亚"] else 0,
    key="origin",
    help="中国工厂按人民币÷汇率；越南/泰国/马来西亚按美元÷1.1含退税",
)

zone = st.sidebar.selectbox("配送区间", [3, 4, 5], index=[3, 4, 5].index(defaults["zone"]) if defaults["zone"] in [3, 4, 5] else 2, key="zone", help="产品开发阶段默认Zone5")

box_count = st.sidebar.number_input("箱数", min_value=1, max_value=10, value=defaults["box_count"], key="box_count")

price = st.sidebar.number_input("售价 ($)", min_value=0.01, value=defaults["price"], step=10.0, key="price")

purchase_price = st.sidebar.number_input(
    "采购成本",
    min_value=0.01,
    value=defaults["purchase"],
    step=50.0,
    key="purchase",
    help="中国工厂填人民币；海外工厂填美元",
)

exchange_rate = st.sidebar.number_input("汇率", min_value=1.0, value=defaults["rate"], step=0.1, key="rate")

st.sidebar.markdown("---")
st.sidebar.markdown("### 包装数据")

col1, col2 = st.sidebar.columns(2)
with col1:
    length_cm = st.number_input("长 (cm)", min_value=1.0, value=defaults["length"], step=1.0, key="length")
with col2:
    width_cm = st.number_input("宽 (cm)", min_value=1.0, value=defaults["width"], step=1.0, key="width")

col3, col4 = st.sidebar.columns(2)
with col3:
    height_cm = st.number_input("高 (cm)", min_value=1.0, value=defaults["height"], step=1.0, key="height")
with col4:
    weight_kg = st.number_input("重量 (kg)", min_value=0.1, value=defaults["weight"], step=1.0, key="weight")


# Build ProductInput
product = ProductInput(
    name=product_name,
    origin=origin,
    zone=zone,
    box_count=box_count,
    price=price,
    purchase_price_cny=purchase_price,
    length_cm=length_cm,
    width_cm=width_cm,
    height_cm=height_cm,
    weight_kg=weight_kg,
    exchange_rate=exchange_rate,
)

# ============================================================
# Business Type Selector
# ============================================================
biz_type = st.radio(
    "选择业务类型",
    ["AMZ FBM (亚马逊自发)", "AMZ FBA (亚马逊代发)", "Walmart WF (出仓价核价)"],
    horizontal=True,
    key="biz_type",
)

st.markdown("---")

# ============================================================
# AMZ FBM Results
# ============================================================
if "FBM" in biz_type:
    st.markdown("## AMZ FBM 利润核算 (6大海外仓自动比价)")

    result = calc_amz_fbm_profit(product)

    # KPI cards
    col1, col2, col3, col4 = st.columns(4)

    with col1:
        margin_color = "positive" if result["margin_off_season"] > 15 else "negative"
        st.markdown(
            f"""
            <div class="metric-card">
                <h3>淡季毛利率</h3>
                <div class="value {margin_color}">{result['margin_off_season']}%</div>
            </div>
            """,
            unsafe_allow_html=True,
        )

    with col2:
        margin_color = "positive" if result["margin_peak_season"] > 15 else "negative"
        st.markdown(
            f"""
            <div class="metric-card">
                <h3>旺季毛利率</h3>
                <div class="value {margin_color}">{result['margin_peak_season']}%</div>
            </div>
            """,
            unsafe_allow_html=True,
        )

    with col3:
        st.markdown(
            f"""
            <div class="metric-card">
                <h3>淡季最优仓</h3>
                <div class="value neutral">{result['best_warehouse_off']}</div>
            </div>
            """,
            unsafe_allow_html=True,
        )

    with col4:
        st.markdown(
            f"""
            <div class="metric-card">
                <h3>旺季最优仓</h3>
                <div class="value neutral">{result['best_warehouse_peak']}</div>
            </div>
            """,
            unsafe_allow_html=True,
        )

    # Cost breakdown table
    st.markdown("### 费用构成明细")

    cost_data = {
        "费用项": ["商品成本", "头程", "尾程物流", "亚马逊佣金", "广告费用", "退货费", "仓储费", "合计"],
        "淡季 ($)": [
            result["cost_off_season"].purchase_cost,
            result["cost_off_season"].shipping_cost,
            result["cost_off_season"].logistics_cost,
            result["cost_off_season"].commission,
            result["cost_off_season"].ad_cost,
            result["cost_off_season"].return_cost,
            result["cost_off_season"].storage_cost,
            result["cost_off_season"].total,
        ],
        "淡季占比 (%)": [
            round(result["cost_off_season"].purchase_cost / result["cost_off_season"].total * 100, 1),
            round(result["cost_off_season"].shipping_cost / result["cost_off_season"].total * 100, 1),
            round(result["cost_off_season"].logistics_cost / result["cost_off_season"].total * 100, 1),
            round(result["cost_off_season"].commission / result["cost_off_season"].total * 100, 1),
            round(result["cost_off_season"].ad_cost / result["cost_off_season"].total * 100, 1),
            round(result["cost_off_season"].return_cost / result["cost_off_season"].total * 100, 1),
            round(result["cost_off_season"].storage_cost / result["cost_off_season"].total * 100, 1),
            100.0,
        ],
        "旺季 ($)": [
            result["cost_peak_season"].purchase_cost,
            result["cost_peak_season"].shipping_cost,
            result["cost_peak_season"].logistics_cost,
            result["cost_peak_season"].commission,
            result["cost_peak_season"].ad_cost,
            result["cost_peak_season"].return_cost,
            result["cost_peak_season"].storage_cost,
            result["cost_peak_season"].total,
        ],
    }
    df_cost = pd.DataFrame(cost_data)
    st.dataframe(df_cost, use_container_width=True, hide_index=True)

    col_a, col_b = st.columns(2)

    with col_a:
        st.markdown("### 淡季利润")
        st.metric("毛利额", f"${result['profit_off_season']:.2f}", f"毛利率 {result['margin_off_season']}%")
        st.metric("总成本", f"${result['cost_off_season'].total:.2f}")

    with col_b:
        st.markdown("### 旺季利润")
        st.metric("毛利额", f"${result['profit_peak_season']:.2f}", f"毛利率 {result['margin_peak_season']}%")
        st.metric("总成本", f"${result['cost_peak_season'].total:.2f}")

    # 6-warehouse comparison
    st.markdown("### 6大海外仓物流费对比")

    wh_data = []
    for wh in WAREHOUSES:
        wh_data.append(
            {
                "海外仓": wh,
                "淡季物流费 ($)": result["warehouse_comparison_off"][wh]["total"],
                "旺季物流费 ($)": result["warehouse_comparison_peak"][wh]["total"],
                "淡季基础运费": result["warehouse_comparison_off"][wh]["base_rate"],
                "淡季附加费": result["warehouse_comparison_off"][wh]["residential"]
                + result["warehouse_comparison_off"][wh]["ahs_dimension"]
                + result["warehouse_comparison_off"][wh]["ahs_weight"]
                + result["warehouse_comparison_off"][wh]["oversize"],
                "旺季附加费": result["warehouse_comparison_peak"][wh]["peak_surcharge"],
                "计费重 (lb)": result["warehouse_comparison_off"][wh]["chargeable_weight"],
            }
        )

    df_wh = pd.DataFrame(wh_data)
    st.dataframe(df_wh, use_container_width=True, hide_index=True)

    fig = go.Figure()
    fig.add_trace(
        go.Bar(
            name="淡季",
            x=WAREHOUSES,
            y=[result["warehouse_comparison_off"][w]["total"] for w in WAREHOUSES],
            marker_color="#6c4ee0",
        )
    )
    fig.add_trace(
        go.Bar(
            name="旺季",
            x=WAREHOUSES,
            y=[result["warehouse_comparison_peak"][w]["total"] for w in WAREHOUSES],
            marker_color="#e8723c",
        )
    )
    fig.update_layout(
        title="各海外仓物流费对比 (淡季 vs 旺季)",
        xaxis_title="海外仓",
        yaxis_title="物流费 ($)",
        barmode="group",
        height=400,
    )
    st.plotly_chart(fig, use_container_width=True)

    # Cost structure pie chart
    st.markdown("### 成本结构分析")
    col_pie1, col_pie2 = st.columns(2)

    with col_pie1:
        labels = ["商品成本", "头程", "尾程物流", "佣金", "广告", "退货", "仓储"]
        values_off = [
            result["cost_off_season"].purchase_cost,
            result["cost_off_season"].shipping_cost,
            result["cost_off_season"].logistics_cost,
            result["cost_off_season"].commission,
            result["cost_off_season"].ad_cost,
            result["cost_off_season"].return_cost,
            result["cost_off_season"].storage_cost,
        ]
        fig_pie = go.Figure(data=[go.Pie(labels=labels, values=values_off, hole=0.4)])
        fig_pie.update_layout(title="淡季成本结构", height=350)
        st.plotly_chart(fig_pie, use_container_width=True)

    with col_pie2:
        values_peak = [
            result["cost_peak_season"].purchase_cost,
            result["cost_peak_season"].shipping_cost,
            result["cost_peak_season"].logistics_cost,
            result["cost_peak_season"].commission,
            result["cost_peak_season"].ad_cost,
            result["cost_peak_season"].return_cost,
            result["cost_peak_season"].storage_cost,
        ]
        fig_pie2 = go.Figure(data=[go.Pie(labels=labels, values=values_peak, hole=0.4)])
        fig_pie2.update_layout(title="旺季成本结构", height=350)
        st.plotly_chart(fig_pie2, use_container_width=True)

    # Shipping info
    st.markdown("### 头程计算详情")
    ship = result["shipping_info"]
    ship_info = pd.DataFrame(
        [
            {"项目": "单箱体积 (m³)", "值": ship["box_volume_m3"]},
            {"项目": "单款总体积 (m³)", "值": ship["total_volume_m3"]},
            {"项目": "单款总重量 (kg)", "值": ship["total_weight_kg"]},
            {"项目": "体积装柜量", "值": round(ship["container_qty"], 0)},
            {"项目": "总头程费用 ($)", "值": ship["total_head_cost"]},
            {"项目": "单款头程费用 ($)", "值": ship["per_unit_head_cost"]},
        ]
    )
    st.dataframe(ship_info, use_container_width=True, hide_index=True)

# ============================================================
# AMZ FBA Results
# ============================================================
elif "FBA" in biz_type:
    st.markdown("## AMZ FBA 利润测算 (中国工厂)")

    result = calc_amz_fba_profit(product)

    col1, col2, col3 = st.columns(3)

    with col1:
        margin_color = "positive" if result["margin"] > 15 else "negative"
        st.markdown(
            f"""
            <div class="metric-card">
                <h3>毛利率</h3>
                <div class="value {margin_color}">{result['margin']}%</div>
            </div>
            """,
            unsafe_allow_html=True,
        )

    with col2:
        st.markdown(
            f"""
            <div class="metric-card">
                <h3>FBA尺寸分段</h3>
                <div class="value neutral" style="font-size:18px;">{result['package_size']}</div>
            </div>
            """,
            unsafe_allow_html=True,
        )

    with col3:
        st.markdown(
            f"""
            <div class="metric-card">
                <h3>计费重 (lb)</h3>
                <div class="value neutral">{result['chargeable_weight']}</div>
            </div>
            """,
            unsafe_allow_html=True,
        )

    st.markdown("---")
    st.markdown("### 费用构成")

    cost_data = {
        "费用项": ["商品成本", "头程 (AGL散货)", "FBA物流费", "亚马逊佣金", "广告费用", "退货费", "仓储费", "合计"],
        "金额 ($)": [
            result["purchase_cost"],
            result["shipping_cost"],
            result["fba_fee"],
            result["commission"],
            result["ad_cost"],
            result["return_cost"],
            result["storage_cost"],
            result["total_cost"],
        ],
        "占比 (%)": [
            round(result["purchase_cost"] / result["total_cost"] * 100, 1),
            round(result["shipping_cost"] / result["total_cost"] * 100, 1),
            round(result["fba_fee"] / result["total_cost"] * 100, 1),
            round(result["commission"] / result["total_cost"] * 100, 1),
            round(result["ad_cost"] / result["total_cost"] * 100, 1),
            round(result["return_cost"] / result["total_cost"] * 100, 1),
            round(result["storage_cost"] / result["total_cost"] * 100, 1),
            100.0,
        ],
    }
    df_cost = pd.DataFrame(cost_data)
    st.dataframe(df_cost, use_container_width=True, hide_index=True)

    col_a, col_b = st.columns(2)
    with col_a:
        st.metric("毛利额", f"${result['profit']:.2f}")
    with col_b:
        st.metric("总成本", f"${result['total_cost']:.2f}")

    # Pie chart
    st.markdown("### 成本结构")
    labels = ["商品成本", "头程", "FBA物流费", "佣金", "广告", "退货", "仓储"]
    values = [
        result["purchase_cost"],
        result["shipping_cost"],
        result["fba_fee"],
        result["commission"],
        result["ad_cost"],
        result["return_cost"],
        result["storage_cost"],
    ]
    fig = go.Figure(data=[go.Pie(labels=labels, values=values, hole=0.4)])
    fig.update_layout(height=400)
    st.plotly_chart(fig, use_container_width=True)

    st.info(
        f"FBA费率说明: 广告费率10% (比FBM高2%) · 退货费率3.5% (比FBM高0.5%) · "
        f"仓储费率1.5% (比FBM高0.5%) · 实际发货数: {result['actual_shipment']}箱"
    )

# ============================================================
# Walmart WF Results
# ============================================================
else:
    st.markdown("## Walmart WF 核价 (出仓价反推)")

    result = calc_wf_profit(product)

    col1, col2, col3, col4 = st.columns(4)

    with col1:
        st.markdown(
            f"""
            <div class="metric-card">
                <h3>出仓价</h3>
                <div class="value neutral">${result['out_price']:.2f}</div>
            </div>
            """,
            unsafe_allow_html=True,
        )

    with col2:
        st.markdown(
            f"""
            <div class="metric-card">
                <h3>前台售价 (低)</h3>
                <div class="value neutral">${result['front_price_low']:.2f}</div>
            </div>
            """,
            unsafe_allow_html=True,
        )

    with col3:
        st.markdown(
            f"""
            <div class="metric-card">
                <h3>前台售价 (高)</h3>
                <div class="value neutral">${result['front_price_high']:.2f}</div>
            </div>
            """,
            unsafe_allow_html=True,
        )

    with col4:
        st.markdown(
            f"""
            <div class="metric-card">
                <h3>CG仓分类</h3>
                <div class="value neutral" style="font-size:16px;">{result['cg_category']}</div>
            </div>
            """,
            unsafe_allow_html=True,
        )

    st.markdown("---")
    st.markdown("### 费用构成")

    cost_data = {
        "费用项": ["商品成本", "头程", "出库操作费", "合计", "扣点22%后出仓价"],
        "金额 ($)": [
            result["purchase_cost"],
            result["shipping_cost"],
            result["outbound_fee"],
            result["total_cost"],
            result["out_price"],
        ],
    }
    df_cost = pd.DataFrame(cost_data)
    st.dataframe(df_cost, use_container_width=True, hide_index=True)

    st.markdown("### 加价率对应的前台售价")
    rates = [1.7, 1.8, 1.9, 2.0, 2.1]
    prices = [result["out_price"] * r for r in rates]
    df_price = pd.DataFrame({"加价率": rates, "前台售价 ($)": [round(p, 2) for p in prices]})
    st.dataframe(df_price, use_container_width=True, hide_index=True)

    fig = go.Figure()
    fig.add_trace(
        go.Bar(
            x=[f"{r}x" for r in rates],
            y=prices,
            marker_color="#6c4ee0",
            text=[f"${p:.2f}" for p in prices],
            textposition="outside",
        )
    )
    fig.update_layout(
        title="加价率与前台售价对应关系",
        xaxis_title="加价率",
        yaxis_title="前台售价 ($)",
        height=350,
    )
    st.plotly_chart(fig, use_container_width=True)

    st.info(
        f"核价公式: 出仓价 = (采购+头程+出库) ÷ 0.78 (即22%扣点 = 15%利润 + 2%货损 + 1%仓储 + 4%扣点) · "
        f"加价率范围: 1.7x - 2.1x"
    )

    # CG仓费率参考表
    st.markdown("### CG仓拣货和发货价格参考")
    cg_data = {
        "分类": list(__import__("calculator").CG_RATES.keys()),
        "基础费用 ($)": [v["base"] for v in __import__("calculator").CG_RATES.values()],
        "首磅后每磅 ($)": [v["per_lb"] for v in __import__("calculator").CG_RATES.values()],
    }
    st.dataframe(pd.DataFrame(cg_data), use_container_width=True, hide_index=True)


# ============================================================
# Footer
# ============================================================
st.markdown("---")
st.markdown(
    """
    <div style="text-align: center; color: #6b7384; font-size: 13px;">
        基于原始 Excel 利润核算表构建 · 支持 AMZ FBM/FBA/WF 三大业务线 · 6大海外仓自动比价<br>
        费率数据有效期: 2024.1.1-2025.1.19 (旺季) / 2026年Q1-Q3 (淡季) · FedEx折扣: 35%
    </div>
    """,
    unsafe_allow_html=True,
)
